"""文档解析：txt/md/csv/docx/pdf。PDF 需要可选安装 pymupdf。"""
from __future__ import annotations

import csv
import pathlib
from dataclasses import dataclass

SUPPORTED_EXTS = {".txt", ".md", ".markdown", ".csv", ".docx", ".pdf"}


@dataclass
class Document:
    id: str
    title: str
    source_path: str
    ext: str
    text: str
    char_count: int


def _read_text(path: pathlib.Path) -> str:
    for encoding in ("utf-8", "gbk", "utf-16"):
        try:
            return path.read_text(encoding=encoding)
        except (UnicodeDecodeError, LookupError):
            continue
    return path.read_text(encoding="utf-8", errors="ignore")


def _parse_docx(path: pathlib.Path) -> str:
    import docx  # python-docx

    document = docx.Document(str(path))
    parts: list[str] = []
    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if not text:
            continue
        style = (paragraph.style.name or "") if paragraph.style else ""
        if style.startswith("Heading"):
            try:
                level = int(style.split()[-1])
            except ValueError:
                level = 1
            parts.append(f"{'#' * min(level, 6)} {text}")
        else:
            parts.append(text)
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip().replace("\n", " ") for cell in row.cells]
            parts.append(" | ".join(cells))
    return "\n\n".join(parts)


def _parse_pdf(path: pathlib.Path) -> str:
    try:
        import fitz  # PyMuPDF
    except ImportError as exc:
        raise RuntimeError(
            "PDF 解析需要 pymupdf，请执行：python -m pip install pymupdf"
        ) from exc
    pages: list[str] = []
    with fitz.open(str(path)) as pdf:
        for index, page in enumerate(pdf, start=1):
            pages.append(f"【第 {index} 页】\n{page.get_text().strip()}")
    return "\n\n".join(pages)


def _parse_csv(path: pathlib.Path) -> str:
    lines: list[str] = []
    with path.open("r", encoding="utf-8", errors="replace", newline="") as handle:
        for row in csv.reader(handle):
            cells = [cell.strip() for cell in row]
            if cells:
                lines.append(" | ".join(cells))
    return "\n".join(lines)


def load_document(path: pathlib.Path) -> Document:
    ext = path.suffix.lower()
    if ext not in SUPPORTED_EXTS:
        raise ValueError(f"不支持的文档类型: {ext}")
    if ext == ".docx":
        text = _parse_docx(path)
    elif ext == ".pdf":
        text = _parse_pdf(path)
    elif ext == ".csv":
        text = _parse_csv(path)
    else:
        text = _read_text(path)
    text = text.strip()
    return Document(
        id=path.name,
        title=path.stem,
        source_path=str(path),
        ext=ext,
        text=text,
        char_count=len(text),
    )


def load_documents(docs_dir: pathlib.Path) -> list[Document]:
    if not docs_dir.exists():
        raise FileNotFoundError(f"文档目录不存在: {docs_dir}")
    documents: list[Document] = []
    skipped: list[str] = []
    for path in sorted(docs_dir.iterdir()):
        if not path.is_file() or path.name.startswith("."):
            continue
        if path.suffix.lower() not in SUPPORTED_EXTS:
            skipped.append(path.name)
            continue
        documents.append(load_document(path))
    if skipped:
        print(f"跳过不支持的格式: {', '.join(skipped)}")
    return documents
